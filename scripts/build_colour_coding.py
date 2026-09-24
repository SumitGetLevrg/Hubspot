import pandas as pd,re
U='/root/.claude/uploads/1df32d6b-5e1e-5686-a9d6-b7cbb640ca98/'
hs=pd.read_excel(U+'82ba3d4b-hubspot-crm-exports-all-contacts-2026-09-23-2.xlsx',dtype=str).fillna('')
tk=pd.read_excel(U+'76b7ecde-hubspot-crm-exports-upcoming-task-2026-09-23.xlsx',dtype=str).fillna('')
import sys
SRC,PREFIX,OUT=(sys.argv[1:4] if len(sys.argv)>3 else (U+'ddd16ab5-Get_Levrg_-_Sphere_of_Influence.xlsx','SOI','/home/user/Hubspot/SOI_July_August_HubSpot_Colour_Coded.xlsx'))
# optional explicit source tab names (July tab, August tab); default '<PREFIX> July' / '<PREFIX> August'
TABS={'July':sys.argv[4] if len(sys.argv)>5 else PREFIX+' July','August':sys.argv[5] if len(sys.argv)>5 else PREFIX+' August'}
ALIAS={'email':['Work Email','Email','Enriched Email'],'phone':['Mobile Phone','Phone'],
       'li':['LinkedIn Profile URL','Linkedin URL','Linkedin Profile URL','Profile URL'],
       'title':['Job Title','Current Job Title'],'company':['Company','Current Company'],
       'domain':['Company Website','Company Domain','Domain']}
def F(r,k):
    for c in ALIAS[k]:
        v=str(r.get(c,'') or '').strip()
        if v: return v
    return ''
# the utm_campaign HubSpot gives this list's imports; a name-only hit carrying it is trusted
CAMPAIGN_UTM={'SOI':'sphere_of_influence','SD':'speed_dating'}.get(PREFIX,'')
x=pd.ExcelFile(SRC)
soi={}
for s in ['July','August']:
    d=pd.read_excel(x,TABS[s],dtype=str).fillna('')
    d=d[[c for c in d.columns if not c.startswith('Unnamed')]]
    d=d[(d.apply(lambda r:''.join(r).strip(),axis=1)!='')]
    soi[s]=d
def ph(p):
    p=re.sub(r'\D','',p or '')
    return p[-10:] if len(p)>=10 else ''
def nm(s): return re.sub(r'[^a-z]','',(s or '').lower())
emap={};pmap={};nmap={}
for i,r in hs.iterrows():
    for e in [r['Email']]+r['Additional email addresses'].split(';'):
        e=e.strip().lower()
        if e: emap.setdefault(e,set()).add(i)
    p=ph(r['Phone Number'])
    if p: pmap.setdefault(p,set()).add(i)
    k=nm(r['First Name'])+'|'+nm(r['Last Name'])
    if k!='|': nmap.setdefault(k,set()).add(i)
up_tids=set(tk['Record ID'])
up_by_c={}
for _,t in tk.iterrows():
    for cid in re.split(r'[;,]',t['Associated Contact IDs']):
        if cid.strip(): up_by_c.setdefault(cid.strip(),[]).append(t)
BLUE={'new','in-progress','in progress','follow up','connected','qualified for sales discovery call'}
BROWN={'not a fit','cancelled','canceled','unqualified'}
def info(i):
    r=hs.loc[i]; cid=r['Record ID']
    tids=[x.strip() for x in r['Associated Task IDs'].split(';') if x.strip()]
    ups=up_by_c.get(cid,[])
    up_ids={t['Record ID'] for t in ups}|{x for x in tids if x in up_tids}
    prev=[x for x in tids if x not in up_ids]
    ls=r['Lead Status'].strip(); utm=r['utm_campaign'].strip(); ds=r['Current Active Deal Stage'].strip()
    col=''
    if ds and 'closed lost' not in ds.lower(): col='Yellow'
    elif not ds and utm and ls.lower() in BLUE and (up_ids or prev): col='Blue'
    elif not ds and utm and ls.lower() in BROWN and prev: col='Brown'
    nxt=min([t['Due date'] for t in ups],default='')
    return dict(rid=cid,ls=ls,utm=utm,ds=ds,nup=len(up_ids),nprev=len(prev),nxt=nxt,
                upt='; '.join(t['Task Title'].strip() for t in ups),color=col,
                hsname=(r['First Name']+' '+r['Last Name']).strip(),hsemail=r['Email'])
GENERIC={'gmail','yahoo','hotmail','outlook','aol','icloud','live','msn','me','protonmail','comcast','att','sbcglobal','verizon','mac','ymail','googlemail'}
def droot(x):
    x=re.sub(r'^https?://(www\.)?','',(x or '').lower()).split('@')[-1].split('/')[0]
    parts=x.split('.')
    return nm(parts[0]) if parts and parts[0] else ''
def nm2(r):
    fn=[t for t in r['First Name'].split() if nm(t) not in ('dr','mr','ms','mrs')]
    return nm(fn[0] if fn else '')+'|'+nm(re.split(r',',r['Last Name'])[0])
def verified(i,r):
    q=hs.loc[i]; em=q['Email'].lower()
    if em.startswith('missing-email-'): return True
    if CAMPAIGN_UTM and q['utm_campaign'].strip().lower()==CAMPAIGN_UTM: return True
    root=droot(em)
    if not root or root in GENERIC: return False
    if F(r,'domain') and droot(F(r,'domain'))==root: return True
    if F(r,'email') and droot(F(r,'email'))==root: return True
    co=nm(F(r,'company'))
    return len(root)>=4 and len(co)>=4 and (root in co or co in root)
def match(r):
    e=F(r,'email').lower(); p=ph(F(r,'phone'))
    E=emap.get(e,set()) if e else set(); P=pmap.get(p,set()) if p else set()
    if E: prim=max(E,key=lambda i:hs.loc[i,'Create Date']); how='Email'
    elif P: prim=max(P,key=lambda i:hs.loc[i,'Create Date']); how='Phone'
    else:
        N=nmap.get(nm(r['First Name'])+'|'+nm(r['Last Name']),set()) or nmap.get(nm2(r),set())
        if not nm(r['Last Name']) or not N: return None,'',[]
        # a name-only hit is trusted only when something else ties it to this person
        V=[i for i in N if verified(i,r)]
        if len(V)==1: prim=V[0]; how='Name'
        elif len(V)>1: prim=max(V,key=lambda i:hs.loc[i,'Create Date']); how='Name'
        else: return None,'Possible',sorted(N)
        return prim,how,[i for i in V if i!=prim]
    # other records for same person: email/phone hits whose name matches, excluding shared-phone colleagues
    fn=nm(r['First Name'].split()[0]) if r['First Name'] else ''
    others=[i for i in (E|P) if i!=prim and (i in E or nm(hs.loc[i,'First Name']).startswith(fn))]
    return prim,how,others
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
FILL={'Blue':'9BC2E6','Brown':'C8A27C','Yellow':'FFE699'}
HDR=PatternFill('solid',fgColor='404040'); HF=Font(bold=True,color='FFFFFF')
def clean(v):
    v=str(v)
    if re.fullmatch(r'=\+?\d+',v): v=v[1:]
    return v
# cross-tab duplicates (same person in July and August)
def li(r): return re.sub(r'^https?://(www\.)?|/$','',F(r,'li').strip().lower())
def pkey(r):
    fn=nm(r['First Name'].split()[0]) if r['First Name'].strip() else ''
    k={('li',li(r)),('em',F(r,'email').lower()),('ph',ph(F(r,'phone'))+'|'+fn)}
    return {x for x in k if x[1] and not x[1].startswith('|')}
# same person in another row (same LinkedIn URL or email) that matched HubSpot directly lends its match
# to rows that have no email/phone of their own
LEND={}
for _d in soi.values():
    for _,_r in _d.iterrows():
        _m=match(_r)
        if _m[0] is not None and _m[1]!='Name':
            for _k in (('li',li(_r)),('em',F(_r,'email').lower())):
                if _k[1]: LEND.setdefault(_k,_m[0])
def match2(r):
    m=match(r)
    if m[0] is not None and m[1]!='Name': return m
    for k in (('li',li(r)),('em',F(r,'email').lower())):
        if k[1] and k in LEND:
            if m[0] is not None and m[0]!=LEND[k]: return LEND[k],'Other row',[m[0]]
            return LEND[k],'Other row',[]
    return m
def fullkey(r):
    k=pkey(r); m=match2(r)[0]
    if m is not None: k.add(('hs',hs.loc[m,'Record ID']))
    return k
keys={s:[fullkey(r) for _,r in d.iterrows()] for s,d in soi.items()}
HSCOLS=['HubSpot Match','Matched By','HubSpot Record ID','HubSpot Name','HubSpot Email','Lead Status','utm_campaign','Current Active Deal Stage','Upcoming Tasks','Next Task Due','Upcoming Task Title(s)','Previous Tasks','Colour Assigned','Notes']
wb=Workbook(); wb.remove(wb.active)
review=[]; summary={}
for s in ['August','July']:
    d=soi[s]; ws=wb.create_sheet(s); src=list(d.columns)
    hdr=src+HSCOLS; ws.append(hdr)
    for c in range(1,len(hdr)+1):
        ws.cell(1,c).fill=HDR; ws.cell(1,c).font=HF
    cnt={'Blue':0,'Brown':0,'Yellow':0,'':0}
    for n,(idx,r) in enumerate(d.iterrows()):
        prim,how,others=match2(r); notes=[]; reasons=[]
        other=[t for t in soi if t!=s and any(keys[s][n]&k for k in keys[t])]
        dup_in=[j+2 for j,k in enumerate(keys[s]) if j!=n and keys[s][n]&k]
        if prim is None:
            I=dict(rid='',hsname='',hsemail='',ls='',utm='',ds='',nup='',nxt='',upt='',nprev='',color='')
            if how=='Possible':
                cand='; '.join(f"{hs.loc[i,'Record ID']} ({hs.loc[i,'First Name']} {hs.loc[i,'Last Name']}, {hs.loc[i,'Email']}, {hs.loc[i,'Lead Status']})" for i in others); others=[]
                reasons.append(f"Not found in HubSpot by email/phone. Possible name match(es) that could not be verified by company/email domain or campaign: {cand} - please verify")
                how=''
            else: reasons.append('Not found in HubSpot (no email / phone / name match)')
        else:
            I=info(prim)
            if how=='Other row': notes.append('No email/phone match of its own; same LinkedIn URL/email as another row in this workbook that matched HubSpot by email/phone'); how='LinkedIn (other row)'
            if how=='Name': notes.append('Matched on name only, verified by import placeholder / campaign utm / email domain vs company - please verify')
            if how=='Phone' and F(r,'email'): notes.append(f"Source email {F(r,'email')} not in HubSpot; matched by phone")
            if not I['color']:
                why=[]
                if I['ds'] and 'closed lost' in I['ds'].lower(): why.append(f"deal stage is {I['ds']}")
                if not I['ds']:
                    if not I['utm']: why.append('utm_campaign is blank')
                    if not (I['nup'] or I['nprev']): why.append('no upcoming or previous tasks')
                    if I['ls'].lower() not in BLUE|BROWN: why.append(f"lead status '{I['ls']}' is not in the Blue or Brown lists")
                    elif I['ls'].lower() in BROWN and not I['nprev']: why.append('Brown lead status but no previous tasks')
                reasons.append('In HubSpot but does not meet colour criteria: '+'; '.join(why))
        if others:
            reasons.append('Duplicate in HubSpot: also matches '+'; '.join(f"{hs.loc[i,'Record ID']} ({hs.loc[i,'First Name']} {hs.loc[i,'Last Name']}, {hs.loc[i,'Email']}, {hs.loc[i,'Lead Status']}{', '+hs.loc[i,'Current Active Deal Stage'] if hs.loc[i,'Current Active Deal Stage'] else ''})" for i in others))
        if other: reasons.append(f"Duplicate: same contact also appears in the {' & '.join(t for t in other)} tab")
        if dup_in: reasons.append('Duplicate: same contact appears more than once in this tab (rows '+', '.join(map(str,dup_in))+')')
        allnotes='; '.join(notes+[x for x in reasons if x.startswith('Duplicate')])
        row=[clean(r[c]) for c in src]+['Yes' if prim is not None else 'No',how,I['rid'],I['hsname'],I['hsemail'],I['ls'],I['utm'],I['ds'],I['nup'],I['nxt'],I['upt'],I['nprev'],I['color'] or 'None',allnotes]
        ws.append(row); cnt[I['color']]+=1
        if I['color']:
            f=PatternFill('solid',fgColor=FILL[I['color']])
            for c in range(1,len(hdr)+1): ws.cell(ws.max_row,c).fill=f
        if reasons:
            review.append([s,n+2,r['First Name'],r['Last Name'],F(r,'title'),F(r,'email'),clean(F(r,'phone')),F(r,'li'),F(r,'company'),I['rid'],I['ls'],I['utm'],I['ds'],I['nup'],I['nprev'],I['color'] or 'None',' | '.join(reasons)])
    summary[s]=cnt
    for c in range(1,len(hdr)+1):
        h=hdr[c-1]; ws.column_dimensions[get_column_letter(c)].width=60 if h in('AI Research','Headline','Notes','Upcoming Task Title(s)') else 22
    for row in ws.iter_rows(min_row=2):
        for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=False)
    ws.freeze_panes='C2'
ws=wb.create_sheet('Not Matched & Duplicates')
H=['Source Tab','Row # in Tab','First Name','Last Name','Job Title','Work Email','Mobile Phone','LinkedIn Profile URL','Company','HubSpot Record ID','Lead Status','utm_campaign','Current Active Deal Stage','Upcoming Tasks','Previous Tasks','Colour Assigned','Reason']
ws.append(H)
for c in range(1,len(H)+1): ws.cell(1,c).fill=HDR; ws.cell(1,c).font=HF; ws.column_dimensions[get_column_letter(c)].width=20
ws.column_dimensions[get_column_letter(len(H))].width=120
for r in review:
    ws.append(r)
    if r[15]!='None':
        for c in range(1,len(H)+1): ws.cell(ws.max_row,c).fill=PatternFill('solid',fgColor=FILL[r[15]])
ws.freeze_panes='C2'
lg=wb.create_sheet('Legend')
L=[['Colour','Rule applied'],
 ['Blue','utm_campaign not blank AND deal stage blank AND lead status in {New, In-Progress, Follow Up, Connected, Qualified for Sales Discovery Call} AND contact has an upcoming task (from the Upcoming Tasks export) or, failing that, a previous task'],
 ['Brown','utm_campaign not blank AND deal stage blank AND lead status in {Not a Fit, Cancelled, Unqualified} AND contact has a previous task'],
 ['Yellow','Current Active Deal Stage is populated and is not a Closed Lost stage'],
 ['None','No colour: not in HubSpot, or in HubSpot but fails the rules above (see the Not Matched & Duplicates tab for the reason)'],
 [],['Matching','Source Work Email vs HubSpot Email + Additional email addresses; then Mobile Phone (last 10 digits); then the match of another row in the workbook with the same LinkedIn URL/email; then exact first+last name, accepted only when verified (HubSpot record is the missing-email import placeholder, carries this list\'s campaign utm_campaign, or its email domain matches the contact\'s company/domain); unverified name hits are listed as possible matches, uncoloured'],
 ['Upcoming task','Contact ID appears in the Upcoming Tasks export'],['Previous task','Task IDs on the HubSpot contact that are not in the Upcoming Tasks export'],
 [],['Tab','Blue','Brown','Yellow','No colour']]
for s,c in summary.items(): L.append([s,c['Blue'],c['Brown'],c['Yellow'],c['']])
for r in L: lg.append(r)
for k,v in FILL.items():
    for row in lg.iter_rows(min_row=2,max_row=4):
        if row[0].value==k: row[0].fill=PatternFill('solid',fgColor=v)
lg.column_dimensions['A'].width=16; lg.column_dimensions['B'].width=140
for c in (lg['A1'],lg['B1'],lg['A7'],lg['A11']): c.font=Font(bold=True)
out=OUT
wb.save(out); print(summary, len(review))

